#!/usr/bin/env python3
"""Local visual proof for the workshop's simple text, shapes, and template artwork."""
import argparse
import json
import os
import subprocess
from io import BytesIO
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.dml import MSO_COLOR_TYPE
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn

DPI = 144
EMU = 914400
FONT_ROOT = Path("/Library/Fonts")
# Template artwork includes SVGs; they are rasterized with the Node "sharp" package.
# Use SHARP_MODULE if set, else build/node_modules/sharp, else Node's normal resolution.
LOCAL_SHARP = Path(__file__).resolve().parent / "node_modules" / "sharp"
SHARP_MODULE = os.environ.get("SHARP_MODULE") or (str(LOCAL_SHARP) if LOCAL_SHARP.exists() else "sharp")


def color(value, default):
    if value.type == MSO_COLOR_TYPE.RGB:
        return "#" + str(value.rgb)
    return default


def font_for(run, size):
    name = run.font.name or "Mona Sans"
    stem = "MonaSansDisplay" if "Display" in name else "MonaSans"
    weight = "Bold" if run.font.bold else "Regular"
    return ImageFont.truetype(str(FONT_ROOT / f"{stem}-{weight}.ttf"), size)


def render(deck, output, only_slide=None):
    prs = Presentation(deck)
    dark = "dark" in Path(deck).stem
    bg, ink = ("#000000", "#FFFFFF") if dark else ("#FFFFFF", "#000000")
    width, height = round(prs.slide_width / EMU * DPI), round(prs.slide_height / EMU * DPI)
    output.mkdir(parents=True, exist_ok=True)
    findings = []
    previews = []
    for number, slide in enumerate(prs.slides, 1):
        if only_slide is not None and number != only_slide:
            continue
        image = Image.new("RGB", (width, height), bg)
        draw = ImageDraw.Draw(image)

        def shape(s, inherited=False):
            if s.is_placeholder or s.shape_type == MSO_SHAPE_TYPE.GROUP:
                return
            x, y, w, h = [round(v / EMU * DPI) for v in (s.left, s.top, s.width, s.height)]
            if x >= width or y >= height:
                return
            if s.shape_type == MSO_SHAPE_TYPE.PICTURE:
                blip = s._element.xpath("./p:blipFill/a:blip")[0]
                rid = blip.get(qn("r:embed"))
                if rid is None:
                    rid = blip.xpath(".//*[local-name()='svgBlip']")[0].get(qn("r:embed"))
                blob = s.part.related_part(rid).blob
                if blob.lstrip().startswith(b"<"):
                    blob = subprocess.run([
                        "node", "-e",
                        "const sharp=require(" + json.dumps(SHARP_MODULE) + ");"
                        "let b=[];process.stdin.on('data',d=>b.push(d));"
                        "process.stdin.on('end',()=>sharp(Buffer.concat(b)).png().toBuffer()"
                        ".then(d=>process.stdout.write(d)));",
                    ], input=blob, capture_output=True, check=True).stdout
                source = Image.open(BytesIO(blob)).convert("RGBA")
                crop = (round(s.crop_left * source.width), round(s.crop_top * source.height),
                        round((1 - s.crop_right) * source.width),
                        round((1 - s.crop_bottom) * source.height))
                source = source.crop(crop).resize((w, h), Image.Resampling.LANCZOS)
                image.paste(source, (x, y), source)
                return
            if s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
                fill = color(s.fill.fore_color, bg) if s.fill.type == 1 else None
                line = color(s.line.color, ink) if s.line.fill.type == 1 else None
                stroke = max(1, round((s.line.width or 12700) / EMU * DPI))
                area = (x, y, x + max(w, 1), y + max(h, 1))
                if s.auto_shape_type == MSO_SHAPE.OVAL:
                    draw.ellipse(area, fill=fill, outline=line, width=stroke)
                elif s.auto_shape_type == MSO_SHAPE.ROUNDED_RECTANGLE:
                    draw.rounded_rectangle(area, radius=min(w, h) * .1,
                                           fill=fill, outline=line, width=stroke)
                else:
                    draw.rectangle(area, fill=fill, outline=line, width=stroke)
            if not s.has_text_frame or not s.text:
                return
            tf = s.text_frame
            inner_w = w - (tf.margin_left + tf.margin_right) / EMU * DPI
            inner_h = h - (tf.margin_top + tf.margin_bottom) / EMU * DPI
            rows = []
            for p in tf.paragraphs:
                if not p.runs:
                    continue
                run = p.runs[0]
                point_size = run.font.size.pt if run.font.size else 18
                font = font_for(run, round(point_size * DPI / 72))
                leading = (point_size * 1.16) * DPI / 72
                spacing = (p.space_after.pt if p.space_after else 0) * DPI / 72
                for forced_line in p.text.replace("\v", "\n").split("\n"):
                    line = ""
                    words = forced_line.split(" ")
                    for word in words:
                        test = (line + " " + word).strip()
                        if line and draw.textlength(test, font=font) > inner_w:
                            rows.append((line, font, leading, color(run.font.color, ink), p.alignment))
                            line = word
                        else:
                            line = test
                    rows.append((line, font, leading, color(run.font.color, ink), p.alignment))
                rows.append(("", font, spacing, ink, p.alignment))
            needed = sum(row[2] for row in rows)
            if needed > inner_h + 3 and not inherited:
                findings.append(f"slide {number}: {needed / DPI:.2f}in of text in "
                                f"{inner_h / DPI:.2f}in: {s.text[:75]!r}")
            ty = y + tf.margin_top / EMU * DPI
            if tf.vertical_anchor == MSO_ANCHOR.MIDDLE:
                ty += max(0, (inner_h - needed) / 2)
            for content, font, leading, fill, align in rows:
                tx = x + tf.margin_left / EMU * DPI
                if align == PP_ALIGN.CENTER:
                    tx += (inner_w - draw.textlength(content, font=font)) / 2
                draw.text((tx, ty), content, font=font, fill=fill, anchor="lt")
                ty += leading

        for inherited in slide.slide_layout.slide_master.shapes:
            shape(inherited, True)
        for inherited in slide.slide_layout.shapes:
            shape(inherited, True)
        for authored in slide.shapes:
            shape(authored)
        file = output / f"slide-{number:02}.png"
        image.save(file)
        previews.append(image.resize((768, 432), Image.Resampling.LANCZOS))
    for first in range(0, len(previews) if only_slide is None else 0, 6):
        group = previews[first:first + 6]
        sheet = Image.new("RGB", (1568, 3 * 470), "#7C8580")
        sd = ImageDraw.Draw(sheet)
        for j, preview in enumerate(group):
            x, y = 8 + (j % 2) * 784, 26 + (j // 2) * 470
            sheet.paste(preview, (x, y))
            sd.text((x + 3, y - 20), f"Slide {first + j + 1}", fill="white")
        sheet.save(output / f"contact-{first // 6 + 1}.jpg", quality=92)
    print("\n".join(findings) if findings else "No estimated text overflow.")
    print(output)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("deck")
    parser.add_argument("output")
    parser.add_argument("--slide", type=int)
    args = parser.parse_args()
    render(args.deck, Path(args.output), args.slide)
