#!/usr/bin/env python3
"""Copy the speaker notes in run-of-show.md into the workshop decks.

Each slide's notes live in run-of-show.md as a blockquote between
`<!-- speaker-notes slide=N -->` and `<!-- /speaker-notes -->`. A slide may have
several blocks; they are joined in document order. Supported Markdown inside a
block: `**bold**`, inline code (backticks are dropped), `- ` bullets, and fenced
code blocks (one monospace notes paragraph per line).

Only the notes parts of each deck are rewritten; every other file in the .pptx
package is copied byte for byte, so hand edits to the slides are preserved.

    python3 build/sync-notes.py            # update both decks
    python3 build/sync-notes.py --check    # exit 1 if a deck's notes differ
"""
import argparse
import os
import re
import sys
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.opc.oxml import serialize_part_xml

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "run-of-show.md"
DECKS = [ROOT / f"gh-skills-builder-workshop-{theme}.pptx" for theme in ("light", "dark")]
MONO = "Courier New"
START = re.compile(r"^<!-- speaker-notes slide=(\d+) -->$")
END = "<!-- /speaker-notes -->"


class NotesError(Exception):
    pass


def inline_runs(text, lineno):
    """Turn one line of prose into (text, bold, mono) runs."""
    runs, bold = [], False
    for part in re.split(r"(`[^`]+`|\*\*)", text):
        if not part:
            continue
        if part.startswith("`") and part.endswith("`") and len(part) > 2:
            runs.append((part[1:-1], bold, False))
        elif part == "**":
            bold = not bold
        else:
            if "`" in part:
                raise NotesError(f"line {lineno}: use balanced single backticks for inline code")
            if re.search(r"\[[^\]]*\]\(", part):
                raise NotesError(f"line {lineno}: Markdown links are not supported in speaker notes")
            if "<" in part:
                raise NotesError(f"line {lineno}: put '<' inside backticks so Markdown does not treat it as HTML")
            if re.search(r"\*|~~|(?<!\w)_(?=\S).*?_(?!\w)", part):
                raise NotesError(f"line {lineno}: use '**bold**' or inline code, not other emphasis")
            runs.append((part, bold, False))
    if bold:
        raise NotesError(f"line {lineno}: unbalanced '**'")
    return runs


def parse_block(lines):
    """Parse one block's (lineno, line) pairs into a list of notes paragraphs."""
    blocks, prose, bullets, code = [], [], [], None

    def flush():
        if prose:
            blocks.append([inline_runs(" ".join(t for _, t in prose), prose[0][0])])
            prose.clear()
        if bullets:
            blocks.append([[("• ", False, False)] + inline_runs(t, n) for n, t in bullets])
            bullets.clear()

    for lineno, raw in lines:
        if not raw.startswith(">"):
            raise NotesError(f"line {lineno}: every speaker-notes line must start with '>'")
        line = raw[2:] if raw.startswith("> ") else raw[1:]
        if code is not None:
            if line.strip() == "```":
                blocks.append(code)
                code = None
            else:
                code.append([(line, False, True)] if line else [])
        elif line.startswith("```"):
            flush()
            code = []
        elif not line.strip():
            flush()
        elif line.startswith("#"):
            raise NotesError(f"line {lineno}: headings are not supported in speaker notes")
        elif re.match(r"\d+[.)] ", line) or line.startswith(("* ", "+ ")):
            raise NotesError(f"line {lineno}: use '- ' bullets; numbered lists are not supported")
        elif line.startswith("- "):
            if prose:
                flush()
            bullets.append((lineno, line[2:]))
        else:
            if bullets:
                flush()
            prose.append((lineno, line.strip()))
    if code is not None:
        raise NotesError("unclosed code fence in speaker notes")
    flush()

    paragraphs = []
    for i, block in enumerate(blocks):
        if i:
            paragraphs.append([])
        paragraphs.extend(block)
    return paragraphs


def parse(markdown):
    """Return {slide number: [paragraph runs]} from run-of-show.md."""
    notes, current, block = {}, None, []
    for lineno, line in enumerate(markdown.splitlines(), 1):
        match = START.match(line.strip())
        if match:
            if current is not None:
                raise NotesError(f"line {lineno}: speaker-notes block opened before the previous one closed")
            current, block = int(match.group(1)), []
        elif line.strip() == END:
            if current is None:
                raise NotesError(f"line {lineno}: closing marker without an opening marker")
            while block and not block[-1][1].strip():
                block.pop()
            while block and not block[0][1].strip():
                block.pop(0)
            paragraphs = parse_block(block)
            if current in notes:
                notes[current] += [[]] + paragraphs
            else:
                notes[current] = paragraphs
            current = None
        elif current is not None:
            block.append((lineno, line))
    if current is not None:
        raise NotesError(f"speaker-notes block for slide {current} is never closed")
    return notes


def normalize(paragraphs):
    """Merge adjacent runs with the same formatting so comparisons ignore run splits."""
    result = []
    for runs in paragraphs:
        merged = []
        for text, bold, mono in runs:
            if not text:
                continue
            if merged and merged[-1][1:] == (bold, mono):
                merged[-1] = (merged[-1][0] + text, bold, mono)
            else:
                merged.append((text, bold, mono))
        result.append(merged)
    while result and not result[-1]:
        result.pop()
    return result


def read_notes(text_frame):
    return normalize(
        [(r.text, bool(r.font.bold), r.font.name == MONO) for r in p.runs]
        for p in text_frame.paragraphs
    )


def write_notes(text_frame, paragraphs):
    text_frame.clear()
    for i, runs in enumerate(paragraphs):
        paragraph = text_frame.paragraphs[0] if i == 0 else text_frame.add_paragraph()
        for text, bold, mono in runs:
            run = paragraph.add_run()
            run.text = text
            if bold:
                run.font.bold = True
            if mono:
                run.font.name = MONO


def replace_parts(deck, parts):
    """Rewrite only the named zip entries; copy every other entry unchanged."""
    tmp = deck.with_name(deck.name + ".tmp")
    with zipfile.ZipFile(deck) as zin, zipfile.ZipFile(tmp, "w") as zout:
        for info in zin.infolist():
            data = parts.get(info.filename)
            zout.writestr(info, data if data is not None else zin.read(info.filename))
    os.replace(tmp, deck)


def sync(deck, notes, check):
    prs = Presentation(str(deck))
    slides = list(prs.slides)
    expected = set(range(1, len(slides) + 1))
    if set(notes) != expected:
        missing, extra = sorted(expected - set(notes)), sorted(set(notes) - expected)
        raise NotesError(f"{deck.name}: {len(slides)} slides; missing notes for {missing}, unknown slides {extra}")

    changed = {}
    for number, slide in enumerate(slides, 1):
        if not slide.has_notes_slide or slide.notes_slide.notes_text_frame is None:
            raise NotesError(f"{deck.name}: slide {number} has no notes placeholder; add one in PowerPoint")
        frame = slide.notes_slide.notes_text_frame
        want = normalize(notes[number])
        want_text = "\n".join("".join(text for text, _, _ in runs) for runs in want)
        # .runs omits PowerPoint soft breaks and fields, but .text includes them.
        if frame.text == want_text and read_notes(frame) == want:
            continue
        if not check:
            write_notes(frame, want)
        part = slide.notes_slide.part
        changed[str(part.partname).lstrip("/")] = serialize_part_xml(part._element)

    numbers = sorted(int(re.search(r"(\d+)\.xml$", name).group(1)) for name in changed)
    if check:
        print(f"{deck.name}: {'notes differ on ' + str(len(changed)) + ' slide(s)' if changed else 'in sync'}")
    elif changed:
        replace_parts(deck, changed)
        print(f"{deck.name}: updated notes parts {numbers}")
    else:
        print(f"{deck.name}: already in sync")
    return not changed


def main():
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("decks", nargs="*", type=Path, default=DECKS)
    parser.add_argument("--source", type=Path, default=SOURCE)
    parser.add_argument("--check", action="store_true", help="report differences without writing")
    args = parser.parse_args()
    try:
        notes = parse(args.source.read_text(encoding="utf-8"))
        in_sync = [sync(deck, notes, args.check) for deck in args.decks]
    except NotesError as error:
        sys.exit(f"sync-notes: {error}")
    if args.check and not all(in_sync):
        sys.exit(1)


if __name__ == "__main__":
    main()
