#!/usr/bin/env python3
"""Template helpers for build-workshop-v2.py: potx conversion, package pruning, and shape utilities."""
from __future__ import annotations

import posixpath
import os
import shutil
import sys
import zipfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
WORKSHOP = HERE.parent
TEMPLATE = Path(os.environ.get(
    "GH_TEMPLATE_POTX",
    WORKSHOP / "template" / "2026 GitHub Presentation Template.potx",
))
TEMPLATE_PPTX = WORKSHOP / "template" / "template-inspect.pptx"

vendor = HERE / f"vendor{sys.version_info.major}{sys.version_info.minor}"
if not vendor.exists():
    vendor = HERE / "vendor"
sys.path.insert(0, str(vendor))
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN, MSO_AUTO_SIZE
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

FONT_DISPLAY = "Mona Sans Display"
FONT_BODY = "Mona Sans"
FONT_MONO = "Mona Sans Mono"

THEMES = {
    "dark": {
        "master": "slideMaster1.xml", "bg": "000000", "panel": "101411", "panel2": "20262C",
        "text": "FFFFFF", "body": "B0BAC3", "muted": "B0BAC3", "green": "5EEC83",
        "green_fill": "0EBE3D", "green_deep": "0D6836", "purple": "B870FF", "purple_deep": "8534F3",
        "card": "09241A", "outline": "255C3B", "card_text": "FFFFFF",
    },
    "light": {
        "master": "slideMaster4.xml", "bg": "FFFFFF", "panel": "F2F5F3", "panel2": "E4EBE6",
        "text": "000000", "body": "20262C", "muted": "5C6A75", "green": "0D6836",
        "green_fill": "0EBE3D", "green_deep": "0D6836", "purple": "8534F3", "purple_deep": "8534F3",
        "card": "EAF5EC", "outline": "A6D9B5", "card_text": "000000",
    },
}


def convert_potx_to_pptx() -> None:
    """Make a PowerPoint presentation package from the template package."""
    if not TEMPLATE.exists():
        raise SystemExit(
            f"Template not found: {TEMPLATE}\n"
            "Set GH_TEMPLATE_POTX to the 2026 GitHub Presentation Template .potx file."
        )
    TEMPLATE_PPTX.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(TEMPLATE) as src, zipfile.ZipFile(TEMPLATE_PPTX, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(
                    b"application/vnd.openxmlformats-officedocument.presentationml.template.main+xml",
                    b"application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml",
                )
            dst.writestr(item, data)


def prune_template_package(path: Path, master_file: str) -> None:
    """Keep only this deck's selected master/layouts and their referenced assets."""
    import posixpath
    import tempfile
    from defusedxml.minidom import parseString

    scripts = Path(os.environ.get(
        "PPTX_SKILL_SCRIPTS",
        Path.home() / ".copilot/installed-plugins/anthropic-agent-skills/document-skills/skills/pptx/scripts",
    ))
    if not (scripts / "clean.py").exists():
        raise SystemExit(
            f"pptx skill scripts not found: {scripts}\n"
            "Set PPTX_SKILL_SCRIPTS to the folder containing the pptx skill's clean.py."
        )
    sys.path.insert(0, str(scripts))
    from clean import clean_unused_files

    rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    rel_tag = f"{{{rel_ns}}}Relationship"

    def read_dom(root: Path, part: str):
        return parseString((root / part).read_bytes())

    def rels_for(root: Path, source: str):
        return read_dom(root, posixpath.join(posixpath.dirname(source), "_rels", posixpath.basename(source) + ".rels"))

    def resolve(source: str, target: str) -> str:
        return posixpath.normpath(posixpath.join(posixpath.dirname(source), target))

    with tempfile.TemporaryDirectory(prefix="workshop-prune-", dir=os.environ.get("TMPDIR")) as td:
        root = Path(td) / "unpacked"
        root.mkdir()
        with zipfile.ZipFile(path) as archive:
            archive.extractall(root)

        presentation_part = "ppt/presentation.xml"
        pres_dom = read_dom(root, presentation_part)
        pres_rels = rels_for(root, presentation_part)
        master_rel_by_path = {}
        rels_to_drop = []
        for rel in list(pres_rels.getElementsByTagName("Relationship")):
            if rel.getAttribute("Type").endswith("/slideMaster"):
                target = resolve(presentation_part, rel.getAttribute("Target"))
                if target.endswith(master_file):
                    master_rel_by_path[target] = rel.getAttribute("Id")
                else:
                    rels_to_drop.append(rel)
        if len(master_rel_by_path) != 1:
            raise ValueError(f"Expected one selected master {master_file}; found {list(master_rel_by_path)}")
        used_master = next(iter(master_rel_by_path))
        keep_master_rid = next(iter(master_rel_by_path.values()))
        for rel in rels_to_drop:
            rel.parentNode.removeChild(rel)
        for el in list(pres_dom.getElementsByTagName("p:sldMasterId")):
            if el.getAttribute("r:id") != keep_master_rid:
                el.parentNode.removeChild(el)
        (root / presentation_part).write_bytes(pres_dom.toxml(encoding="utf-8"))
        pres_rel_path = posixpath.join(posixpath.dirname(presentation_part), "_rels", posixpath.basename(presentation_part) + ".rels")
        (root / pres_rel_path).write_bytes(pres_rels.toxml(encoding="utf-8"))

        used_layouts = set()
        for slide_path in (root / "ppt/slides").glob("slide*.xml"):
            rel_path = slide_path.parent / "_rels" / f"{slide_path.name}.rels"
            if not rel_path.exists():
                continue
            dom = parseString(rel_path.read_bytes())
            for rel in dom.getElementsByTagName("Relationship"):
                if rel.getAttribute("Type").endswith("/slideLayout"):
                    used_layouts.add(resolve(f"ppt/slides/{slide_path.name}", rel.getAttribute("Target")))
        if not used_layouts:
            raise ValueError("No slide layouts are referenced by output slides")

        # The layout parts identify their owning master; this also guards against pruning the wrong theme.
        layout_masters = set()
        for layout in used_layouts:
            rel_path = root / posixpath.dirname(layout) / "_rels" / (posixpath.basename(layout) + ".rels")
            dom = parseString(rel_path.read_bytes())
            for rel in dom.getElementsByTagName("Relationship"):
                if rel.getAttribute("Type").endswith("/slideMaster"):
                    layout_masters.add(resolve(layout, rel.getAttribute("Target")))
        if layout_masters != {used_master}:
            raise ValueError(f"Slide layouts span unexpected masters: {layout_masters}")

        deleted_parts = set()
        masters_dir = root / "ppt/slideMasters"
        for master_path in masters_dir.glob("slideMaster*.xml"):
            rel_path = master_path.parent / "_rels" / f"{master_path.name}.rels"
            source = f"ppt/slideMasters/{master_path.name}"
            if source != used_master:
                deleted_parts.add(source)
                if rel_path.exists():
                    unused_master_rels = parseString(rel_path.read_bytes())
                    for rel in unused_master_rels.getElementsByTagName("Relationship"):
                        if rel.getAttribute("Type").endswith("/slideLayout"):
                            target = resolve(source, rel.getAttribute("Target"))
                            if target in used_layouts:
                                raise ValueError(f"A selected slide depends on unused master {source}")
                            deleted_parts.add(target)
                            target_rels = root / posixpath.dirname(target) / "_rels" / (posixpath.basename(target) + ".rels")
                            if target_rels.exists():
                                target_rels.unlink()
                                deleted_parts.add(str(target_rels.relative_to(root)).replace("\\", "/"))
                            (root / target).unlink(missing_ok=True)
                    rel_path.unlink()
                    deleted_parts.add(str(rel_path.relative_to(root)).replace("\\", "/"))
                master_path.unlink()
                continue
            master_dom = parseString(master_path.read_bytes())
            master_rels = parseString(rel_path.read_bytes())
            keep_layout_rids = set()
            for rel in list(master_rels.getElementsByTagName("Relationship")):
                if rel.getAttribute("Type").endswith("/slideLayout"):
                    target = resolve(source, rel.getAttribute("Target"))
                    if target in used_layouts:
                        keep_layout_rids.add(rel.getAttribute("Id"))
                    else:
                        rel.parentNode.removeChild(rel)
                        deleted_parts.add(target)
                        target_rels = root / posixpath.dirname(target) / "_rels" / (posixpath.basename(target) + ".rels")
                        if target_rels.exists():
                            target_rels.unlink()
                            deleted_parts.add(str(target_rels.relative_to(root)).replace("\\", "/"))
                        (root / target).unlink(missing_ok=True)
            for el in list(master_dom.getElementsByTagName("p:sldLayoutId")):
                if el.getAttribute("r:id") not in keep_layout_rids:
                    el.parentNode.removeChild(el)
            master_path.write_bytes(master_dom.toxml(encoding="utf-8"))
            rel_path.write_bytes(master_rels.toxml(encoding="utf-8"))

        content_path = root / "[Content_Types].xml"
        content_dom = parseString(content_path.read_bytes())
        for override in list(content_dom.getElementsByTagName("Override")):
            if override.getAttribute("PartName").lstrip("/") in deleted_parts:
                override.parentNode.removeChild(override)
        content_path.write_bytes(content_dom.toxml(encoding="utf-8"))

        clean_unused_files(root)
        replacement = path.with_name(path.stem + "-pruned.pptx")
        with zipfile.ZipFile(replacement, "w", zipfile.ZIP_DEFLATED, compresslevel=6) as archive:
            for file_path in root.rglob("*"):
                if file_path.is_file():
                    archive.write(file_path, file_path.relative_to(root).as_posix())
        shutil.move(str(replacement), str(path))

def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def add_text(slide, x, y, w, h, text, *, size=18, color="000000", bold=False,
             font=FONT_BODY, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0.04,
             fit=True):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    if fit:
        tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    tf.vertical_anchor = valign
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(3)
        run = p.add_run()
        run.text = line
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = rgb(color)
    return shape


def add_box(slide, x, y, w, h, fill, line, *, radius=True, width=1.0):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line)
    shape.line.width = Pt(width)
    return shape


def add_badge(slide, x, y, d, label, t, fill=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill or t["green_fill"])
    shape.line.fill.background()
    tf = shape.text_frame
    tf.clear()
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.name = FONT_DISPLAY
    r.font.size = Pt(16 if d < 0.55 else 22)
    r.font.bold = True
    r.font.color.rgb = rgb("FFFFFF")
    return shape


def layout_by_name(prs, name, master_part):
    for s in prs.slides:
        layout = s.slide_layout
        if layout.name == name and layout.slide_master.part.partname.endswith(master_part):
            return layout
    raise ValueError(f"No layout {name!r} on {master_part}")


def set_notes(slide, text):
    if text:
        slide.notes_slide.notes_text_frame.text = text


if __name__ == "__main__":
    raise SystemExit(
        "build-from-template.py is a helper module for build-workshop-v2.py.\n"
        "Run build/build-workshop-v2.py to build the decks."
    )
