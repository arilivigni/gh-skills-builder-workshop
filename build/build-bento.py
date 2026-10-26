#!/usr/bin/env python3
"""Build the Bento slide decks from the dark PowerPoint deck and run-of-show.md.

Bento (https://bento.page, MIT) decks are single, self-contained HTML files: the
Bento app plus one `<script type="application/bento+json" id="bento-doc">` block
that holds the slides as JSON. Reveal.js drives navigation when presenting.
This script rebuilds that JSON block and nothing else:

- Images come from gh-skills-builder-workshop-dark.pptx, recompressed for the web.
- Speaker notes come from run-of-show.md, parsed by sync-notes.py.
- Fonts are the Mona Sans woff2 files in site/assets/fonts/ (SIL OFL 1.1).

It writes three things:

- gh-skills-builder-workshop-dark.bento.html: the editable deck, with notes.
- site/slides/gh-skills-builder-workshop.bento.html: a read-only player copy
  for the website, without speaker notes.
- site/assets/img/: the same images, for the website.

Edit slides here, not in the Bento editor: the next run overwrites the JSON.

    python3 build/build-bento.py                  # update decks and site images
    python3 build/build-bento.py --check          # exit 1 if anything is out of date
    python3 build/build-bento.py --refresh-shell  # download the latest Bento app first
    python3 build/build-bento.py --shell FILE     # or use a Bento app saved locally
"""
import argparse
import base64
import copy
import html
import importlib.util
import io
import json
import re
import sys
import urllib.request
from datetime import datetime, timezone
from pathlib import Path
from xml.etree import ElementTree

from PIL import Image
from pptx import Presentation

ROOT = Path(__file__).resolve().parent.parent
PPTX = ROOT / "gh-skills-builder-workshop-dark.pptx"
NOTES_SOURCE = ROOT / "run-of-show.md"
DECK = ROOT / "gh-skills-builder-workshop-dark.bento.html"
SITE_DECK = ROOT / "site" / "slides" / "gh-skills-builder-workshop.bento.html"
SITE_IMG = ROOT / "site" / "assets" / "img"
FONTS = ROOT / "site" / "assets" / "fonts"
SHELL_URL = "https://bento.page/releases/slides/Bento_Slides.bento.html"
BLOCK = re.compile(r'(<script type="application/bento\+json" id="bento-doc">)(.*?)(</script>)', re.S)
PLUGIN_URL = "https://gh.io/gh-skills-builder"

BG = "#101411"
PANEL = "#09241A"
BADGE = "#0D6836"
BORDER = "#2A6B3F"
WHITE = "#FFFFFF"
MUTED = "#B0BAC3"
ACCENT = "#5EEC83"
SANS = "'Mona Sans', system-ui, -apple-system, 'Segoe UI', sans-serif"
MONO = "'Mona Sans Mono', ui-monospace, SFMono-Regular, Menlo, monospace"
FONT_FILES = [
    ("Mona Sans", "400", "MonaSans-Regular.woff2"),
    ("Mona Sans", "600", "MonaSans-SemiBold.woff2"),
    ("Mona Sans", "700", "MonaSans-Bold.woff2"),
    ("Mona Sans Mono", "400", "MonaSansMono-Regular.woff2"),
    ("Mona Sans Mono", "600", "MonaSansMono-SemiBold.woff2"),
]


class BuildError(Exception):
    pass


def load_sync_notes():
    spec = importlib.util.spec_from_file_location("sync_notes", Path(__file__).with_name("sync-notes.py"))
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


# --- Media ---------------------------------------------------------------


def encode(img, fmt, **options):
    buf = io.BytesIO()
    img.save(buf, fmt, **options)
    return buf.getvalue()


def fit_width(img, width):
    if img.width <= width:
        return img
    return img.resize((width, round(img.height * width / img.width)), Image.LANCZOS)


def clean_svg(markup):
    """Drop Office theme classes so the mark renders white anywhere, as an <img> or inline."""
    # Office writes CRLF; Git would store LF and --check would then flag the SVGs.
    markup = markup.replace("\r\n", "\n").replace("\r", "\n")
    markup = re.sub(r"<style>.*?</style>", "", markup, flags=re.S)

    def unclass(match):
        tag = match.group(0).replace(' class="MsftOfcThm_Text1_Fill_v2"', "")
        return tag if re.search(r'\sfill="', tag) else re.sub(r"^<([\w:]+)", r'<\1 fill="#FFFFFF"', tag)

    markup = re.sub(r'<[\w:]+\s[^>]*class="MsftOfcThm_Text1_Fill_v2"[^>]*>', unclass, markup)
    if "class=" in markup:
        raise BuildError("unexpected class attribute in the GitHub mark SVG")
    try:
        ElementTree.fromstring(markup)
    except ElementTree.ParseError as error:
        raise BuildError(f"the GitHub mark SVG is not well-formed XML: {error}") from None
    return markup.strip()


def extract_media(pptx_path):
    """Pull the deck's images out of the PowerPoint file, by slide, and optimize them."""
    prs = Presentation(str(pptx_path))
    slides = list(prs.slides)

    def pictures(number):
        found = {}
        for shape in slides[number - 1].shapes:
            try:
                picture = shape.image
            except (AttributeError, ValueError):
                continue
            found[shape.name + ":" + picture.content_type] = picture
        return found

    def layout_part(number, content_type):
        layout = slides[number - 1].slide_layout
        parts = [r.target_part for r in layout.part.rels.values()
                 if "image" in r.reltype and r.target_part.content_type == content_type]
        if not parts:
            raise BuildError(f"slide {number} layout has no {content_type} image")
        return parts[0].blob

    def only_picture(number):
        found = list(pictures(number).values())
        if len(found) != 1:
            raise BuildError(f"expected one picture on slide {number}, found {len(found)}")
        return Image.open(io.BytesIO(found[0].blob))

    cover = pictures(1)
    try:
        rail = Image.open(io.BytesIO(cover["Picture 10:image/png"].blob))
        photo = Image.open(io.BytesIO(cover["Picture Placeholder 2:image/jpeg"].blob)).convert("RGB")
    except KeyError as missing:
        raise BuildError(f"slide 1 is missing picture {missing}") from None

    # The cover photo frame is 129 x 89 slide pixels: crop to that shape, centered.
    target = 129 / 89
    crop_w = round(photo.height * target)
    left = (photo.width - crop_w) // 2
    photo = photo.crop((left, 0, left + crop_w, photo.height)).resize((387, 267), Image.LANCZOS)

    qr_images = list(pictures(4).values())
    if len(qr_images) != 1 or qr_images[0].content_type != "image/png":
        raise BuildError("expected one PNG QR code on slide 4")

    texture = Image.open(io.BytesIO(layout_part(15, "image/png"))).convert("RGB")
    mark = clean_svg(layout_part(2, "image/svg+xml").decode("utf-8"))

    media = {
        "rail": ("image/webp", encode(fit_width(rail.convert("RGBA"), 470), "WEBP", quality=90, method=6)),
        "headshot": ("image/webp", encode(photo, "WEBP", quality=85, method=6)),
        "qr": ("image/png", qr_images[0].blob),
        "texture": ("image/webp", encode(fit_width(texture, 1920), "WEBP", quality=72, method=6)),
        "art-agent-skills": ("image/webp", encode(fit_width(only_picture(7).convert("RGB"), 1920), "WEBP", quality=80, method=6)),
        "art-lifecycle": ("image/webp", encode(fit_width(only_picture(8).convert("RGB"), 1920), "WEBP", quality=80, method=6)),
        "art-review-circuit": ("image/webp", encode(fit_width(only_picture(11).convert("RGB"), 1920), "WEBP", quality=80, method=6)),
        "github-mark": ("image/svg+xml", mark.encode("utf-8")),
    }

    # Open Graph preview: the lifecycle art, cropped to 1200 x 630.
    og = fit_width(only_picture(8).convert("RGB"), 1200)
    top = (og.height - 630) // 2
    media["og-image"] = ("image/jpeg", encode(og.crop((0, top, 1200, top + 630)), "JPEG", quality=85, optimize=True))
    return media


def site_images(media):
    """Files for site/assets/img/, keyed by file name."""
    ext = {"image/webp": "webp", "image/png": "png", "image/jpeg": "jpg", "image/svg+xml": "svg"}
    files = {f"{key}.{ext[mime]}": data for key, (mime, data) in media.items() if key != "rail"}
    mark_paths = re.search(r"<svg[^>]*>(.*)</svg>", media["github-mark"][1].decode("utf-8"), re.S).group(1)
    files["favicon.svg"] = (
        '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 256 256">'
        f'<rect width="256" height="256" rx="56" fill="{BG}"/>'
        f'<g transform="translate(27 30.5)">{mark_paths.replace("#FFFFFF", ACCENT)}</g></svg>\n'
    ).encode("utf-8")
    return files


def data_uri(mime, data):
    return f"data:{mime};base64,{base64.b64encode(data).decode('ascii')}"


# --- Elements --------------------------------------------------------------


def esc(*lines):
    return "<br>".join(html.escape(line, quote=False) for line in lines)


def frame(el_id, x, y, w, h, opacity=1, fx=None, link=None, role=None, rotation=0):
    el = {"id": el_id, "x": x, "y": y, "w": w, "h": h, "rotation": rotation, "opacity": opacity}
    if fx:
        el["fx"] = fx
    if link:
        el["link"] = link
    if role:
        el["role"] = role
    return el


def text(el_id, x, y, w, h, markup, size, weight=400, color=WHITE, family=SANS, align="left",
         valign="top", line=1.25, letter=None, **kw):
    el = frame(el_id, x, y, w, h, **kw)
    el.update(type="text", html=markup, fontSize=size, fontFamily=family, fontWeight=weight,
              color=color, align=align, valign=valign, lineHeight=line)
    if letter is not None:
        el["letterSpacing"] = letter
    return el


def shape(el_id, kind, x, y, w, h, fill, stroke="transparent", stroke_width=0, radius=0,
          stroke_style=None, **kw):
    el = frame(el_id, x, y, w, h, **kw)
    el.update(type="shape", shape=kind, fill=fill, stroke=stroke, strokeWidth=stroke_width, radius=radius)
    if stroke_style:
        el["strokeStyle"] = stroke_style
    return el


def image(el_id, x, y, w, h, asset, fit="cover", radius=0, **kw):
    el = frame(el_id, x, y, w, h, **kw)
    el.update(type="image", src=f"asset:{asset}", fit=fit, radius=radius)
    return el


def svg(el_id, x, y, w, h, asset, **kw):
    el = frame(el_id, x, y, w, h, **kw)
    el.update(type="svg", asset=asset)
    return el


def enter(order, kind="fade-up"):
    return {"enter": kind, "order": order}


KEN = {"ambient": "kenburns", "ken": {"dir": "drift", "scale": 1.04, "duration": 30}}


def chrome(fx=None):
    """The bottom rule and GitHub mark from the PowerPoint layout. Shared ids keep it still."""
    return [
        shape("rule", "rect", 0, 676, 1280, 2, ACCENT, opacity=0.55, fx=fx),
        shape("mark-divider", "rect", 1236, 678, 1, 42, "#3A4440", fx=fx),
        svg("mark", 1246, 687, 24, 23, "github-mark", fx=fx),
    ]


def heading(prefix, title, subtitle=None):
    els = [text(f"{prefix}-title", 96, 60, 1088, 56, esc(title), 44, 700, line=1.2, role="title", fx=enter(0))]
    if subtitle:
        els.append(text(f"{prefix}-subtitle", 96, 128, 1088, 36, esc(subtitle), 24, color=MUTED, line=1.3,
                        role="subtitle", fx=enter(0)))
    return els


def footnote(prefix, words):
    return text(f"{prefix}-footnote", 96, 628, 1088, 24, esc(words), 15, color=MUTED, line=1.3, fx=enter(9, "fade"))


def badge(el_id, label_id, x, y, size, label, order, font=None):
    return [
        shape(el_id, "ellipse", x, y, size, size, BADGE, fx=enter(order)),
        text(label_id, x, y, size, size, esc(label), font or round(size * 0.42), 700, align="center",
             valign="middle", line=1, fx=enter(order)),
    ]


def card(prefix, card_id, badge_id, x, y, w, h, mark, title, body, order, title_size=28, body_size=21,
         fill="transparent"):
    pad = 24
    return [
        shape(card_id, "rect", x, y, w, h, fill, BORDER, 1.5, 28, fx=enter(order)),
        *badge(badge_id, f"{prefix}-{badge_id}-label", x + pad, y + pad, 44, mark, order),
        text(f"{prefix}-{card_id}-title", x + pad, y + 92, w - 2 * pad, 40, esc(title), title_size, 700,
             line=1.2, fx=enter(order)),
        text(f"{prefix}-{card_id}-body", x + pad, y + 150, w - 2 * pad, h - 174, esc(*body), body_size,
             color=MUTED, line=1.4, role="body", fx=enter(order)),
    ]


PHASES = ["Outline", "Bootstrap", "Review", "Publish"]


def tracker(active, y=598):
    """Four phase pills; the highlight and label colors morph from slide to slide."""
    els = []
    for i in range(4):
        els.append(shape(f"phase-pill-{i}", "rect", 744 + i * 112, y, 104, 32, BG, BORDER, 1.5, 16))
    els.append(shape("phase-active", "rect", 744 + active * 112, y, 104, 32, BADGE, ACCENT, 1.5, 16))
    for i, phase in enumerate(PHASES):
        els.append(text(f"phase-label-{i}", 744 + i * 112, y, 104, 32, esc(phase), 14, 600,
                        WHITE if i == active else MUTED, align="center", valign="middle", line=1))
    return els


def command(prefix, words):
    width = round(len(words) * 13.2) + 40
    return [
        shape("command-chip", "rect", 96, 592, width, 44, PANEL, BORDER, 1.5, 10),
        text(f"{prefix}-command", 116, 592, width - 20, 44, esc(words), 22, 600, ACCENT, MONO,
             valign="middle", line=1, fx=enter(8, "fade")),
    ]


def art(prefix, asset):
    return image(f"{prefix}-art", 0, 0, 1280, 720, asset, fx=KEN)


# --- Slides ----------------------------------------------------------------


def slide(slide_id, name, elements, transition="fade", background=BG, hidden=False):
    s = {"id": slide_id, "name": name, "background": background, "transition": transition,
         "elements": elements, "notes": ""}
    if hidden:
        s["hidden"] = True
    return s


def drop_morphing_entrances(slides):
    """An element that morphs from the previous slide is already moving, so Bento skips its entrance."""
    previous = set()
    for s in slides:
        ids = {el.get("morphId", el["id"]) for el in s["elements"]}
        if s["transition"] == "morph":
            for el in s["elements"]:
                if el.get("morphId", el["id"]) in previous and "fx" in el:
                    el["fx"] = {k: v for k, v in el["fx"].items() if k not in ("enter", "order", "countUp")}
                    if not el["fx"]:
                        del el["fx"]
        previous = ids
    return slides


def build_slides():
    agenda = [
        ("00-06", "Get oriented", "Plugins, installation, and the quick-start scenario"),
        ("06-16", "Outline + hand-off", "Three learner steps, a name, and a private repo"),
        ("16-27", "Bootstrap", "From outline to files, plus one fail-to-pass check"),
        ("27-35", "Review", "Report findings, approve one fix, review again"),
        ("35-41", "Publish readiness", "Evidence and a launch plan, without publishing"),
        ("41-60", "Recap + Q&A", "Takeaways by 0:45, then a buffer and Q&A by 0:50"),
    ]
    agenda_rows = []
    for i, (when, what, detail) in enumerate(agenda):
        y = 72 + i * 94
        fx = enter(i + 1)
        agenda_rows += [
            shape(f"agenda-dot-{i}", "ellipse", 438, y + 9, 12, 12, ACCENT, fx=fx),
            text(f"agenda-time-{i}", 470, y, 104, 30, esc(when), 22, 600, ACCENT, MONO, line=1.3, fx=fx),
            text(f"agenda-what-{i}", 586, y - 2, 598, 34, esc(what), 26, 700, line=1.25, fx=fx),
            text(f"agenda-detail-{i}", 586, y + 34, 598, 28, esc(detail), 19, color=MUTED, line=1.35, fx=fx),
        ]

    steps = [
        ("1", "Find it", ["Open Plugins in the Copilot app.", "Find gh-skills-builder."]),
        ("2", "Review + install", ["Review the source and capabilities.", "Install and enable for your session."]),
        ("3", "Use it", ["Choose the orchestrator agent.", "Ask it to run the exercise lifecycle."]),
    ]
    install_steps = []
    for i, (num, title, body) in enumerate(steps):
        y = 216 + i * 104
        install_steps += [
            *badge(f"install-badge-{i}", f"install-badge-{i}-label", 96, y + 1, 36, num, i + 1, 18),
            text(f"install-step-{i}", 148, y, 560, 34, esc(title), 26, 700, line=1.25, fx=enter(i + 1)),
            text(f"install-step-{i}-body", 148, y + 36, 560, 56, esc(*body), 19, color=MUTED, line=1.4,
                 fx=enter(i + 1)),
        ]

    return [
        slide("welcome", "Welcome", [
            image("cover-texture", 235, 0, 1045, 720, "texture", opacity=0.3, fx=KEN | {"ken": {"dir": "drift", "scale": 1.08, "duration": 24}}),
            image("cover-rail", 0, 0, 235, 720, "rail"),
            text("cover-title", 256, 28, 928, 170, esc("Turn any idea into a Hands-on GitHub Exercise"), 66,
                 line=1.08, role="title", fx=enter(0)),
            text("cover-subtitle", 256, 204, 928, 40, esc("GitHub Copilot app/CLI and the gh-skills-builder-plugin"),
                 30, color=ACCENT, line=1.2, role="subtitle", fx=enter(1)),
            shape("cover-divider", "rect", 256, 606, 928, 1, "#5A6560", fx=enter(2, "fade")),
            image("cover-photo", 256, 622, 129, 89, "headshot", fx=enter(2, "fade")),
            text("cover-name", 404, 624, 600, 30, esc("Ari LiVigni"), 24, 600, family=MONO, line=1.2,
                 fx=enter(2, "fade")),
            text("cover-role", 404, 656, 700, 26, esc("Sr. Learning Advocate and Agentic AI Engineer"), 18, 600,
                 ACCENT, line=1.3, fx=enter(2, "fade")),
            text("cover-company", 404, 684, 300, 26, esc("GitHub"), 18, 600, ACCENT, line=1.3, fx=enter(2, "fade")),
        ]),
        slide("flight-plan", "Flight plan", [
            *chrome(),
            text("agenda-title", 96, 72, 88, 52, "Our ", 42, 700, line=1.15, role="title", fx=enter(0)),
            text("agenda-duration", 184, 72, 56, 52, "60", 42, 700, line=1.15,
                 fx=enter(0) | {"countUp": True}),
            text("agenda-unit", 240, 72, 194, 52, "-minute", 42, 700, line=1.15, fx=enter(0)),
            text("agenda-title-rest", 96, 120, 340, 64, "flight plan", 42, 700, line=1.15, fx=enter(0)),
            text("agenda-note", 96, 512, 316, 96, esc("Mostly live.", "Checkpoints protect the demo."), 22,
                 color=MUTED, line=1.4, role="subtitle", fx=enter(7, "fade")),
            # A horizontal line rotated 90 degrees around its center (x 444, y 320).
            shape("agenda-path", "line", 204, 319, 480, 2, ACCENT, ACCENT, 2, stroke_style="dashed", opacity=0.7,
                  rotation=90, fx={"loop": {"type": "dash-march", "distance": 18, "duration": 1.4}}),
            *agenda_rows,
        ]),
        slide("plugin", "What is a plugin?", [
            *chrome(),
            *heading("plugin", "What is a plugin?", "An installable toolkit that adds reusable capabilities to GitHub Copilot."),
            shape("card-a", "rect", 96, 216, 464, 352, PANEL, BORDER, 1.5, 28, fx=enter(1)),
            *badge("badge-a", "plugin-badge-a-label", 128, 248, 64, "+", 1, 28),
            text("plugin-name", 212, 258, 330, 44, esc("gh-skills-builder"), 34, 700, line=1.2, fx=enter(1)),
            text("plugin-pitch", 128, 360, 400, 100, esc("Install once.", "Reuse the workflow."), 34, 700, ACCENT,
                 line=1.35, fx=enter(1)),
            *card("plugin", "card-b", "badge-b", 592, 216, 280, 352, "A", "One agent",
                  ["Coordinates the work and asks for your approval."], 2, 30, 22),
            *card("plugin", "card-c", "badge-c", 904, 216, 280, 352, "S", "Five skills",
                  ["Reusable workflows and capabilities the agent loads."], 3, 30, 22),
            footnote("plugin", "A plugin packages customizations. It is not a separate AI model."),
        ]),
        slide("install", "Install the toolkit", [
            *chrome(),
            *heading("install", "Add the toolkit. Start building.",
                     "gh-skills-builder is available in the Awesome Copilot marketplace."),
            *install_steps,
            image("install-qr", 932, 192, 224, 224, "qr", fit="contain", link=PLUGIN_URL, fx=enter(4, "fade")),
            text("install-link", 904, 428, 280, 32, "<u>gh.io/gh-skills-builder</u>", 22, 600, ACCENT,
                 align="center", line=1.3, link=PLUGIN_URL, fx=enter(4, "fade")),
            shape("install-cli-box", "rect", 96, 560, 1088, 48, "transparent", BORDER, 1.5, 10, fx=enter(5, "fade")),
            text("install-cli", 116, 560, 1048, 48, esc("CLI: copilot plugin install gh-skills-builder@awesome-copilot"),
                 18, family=MONO, valign="middle", line=1.2, fx=enter(5, "fade")),
        ]),
        slide("scenario", "Today's scenario", [
            *chrome(),
            *heading("scenario", "Today's scenario: Tailspin Toys",
                     "Build a 15–20-minute Copilot quick-start inside Tailspin Toys."),
            *card("scenario", "card-a", "badge-a", 96, 216, 496, 296, "1", "Start with a real app",
                  ["Tailspin Toys as MIT starter code. Learners customize Copilot, not app features."], 1, 30, 22),
            *badge("scenario-arrow", "scenario-arrow-label", 620, 344, 40, ">", 2, 20),
            *card("scenario", "card-b", "badge-b", 688, 216, 496, 296, "2", "Teach three habits",
                  ["Repository instructions.", "Consistent quality reports.", "A report-only QA agent."], 3, 30, 22),
            text("scenario-title-vote", 96, 552, 1088, 40,
                 esc("Rehearsed title: Cleared for Takeoff. The room can vote for another."), 28, 700, ACCENT,
                 line=1.3, fx=enter(4)),
        ]),
        slide("architect", "From builder to architect", [
            *chrome(),
            *heading("architect", "From builder to architect",
                     "The agent does the heavy lifting. You own the learning experience."),
            *card("architect", "card-a", "badge-a", 96, 216, 528, 320, "A", "Let the agent build",
                  ["Draft step content.", "Scaffold repository files.", "Wire validation and feedback."], 1, 30, 22),
            *card("architect", "card-b", "badge-b", 656, 216, 528, 320, "H", "Keep the human decisions",
                  ["Set learner outcomes.", "Approve the outline and fixes.", "Authorize publication."], 2, 30, 22),
            footnote("architect",
                     "Same pattern for educators, maintainers, event organizers, DevRel, startups, and enterprises."),
        ], transition="morph"),
        slide("agent-skills", "One agent, five skills", [art("agent-skills", "art-agent-skills")]),
        slide("lifecycle", "From idea to exercise", [art("lifecycle", "art-lifecycle")]),
        slide("outline", "Outline", [
            *chrome(),
            *heading("outline", "Outline: decide what success looks like",
                     "Input: a Tailspin idea. Output: an approved three-step journey."),
            *card("outline", "panel-1", "badge-1", 96, 216, 340, 336, "1", "Outcome",
                  ["Three small Copilot customizations, done in 15–20 minutes."], 1),
            *card("outline", "panel-2", "badge-2", 470, 216, 340, 336, "2", "Identity",
                  ["Choose an exercise title and an eye-catching repository name."], 2),
            *card("outline", "panel-3", "badge-3", 844, 216, 340, 336, "3", "Validation",
                  ["For every graded step: what can a workflow actually check?"], 3),
            *command("outline", "/create-github-skills-outline"),
            *tracker(0),
        ]),
        slide("bootstrap", "Bootstrap", [
            *chrome(),
            *heading("bootstrap", "Bootstrap: make the outline runnable",
                     "Generate exercise files in the current repository, not a new GitHub remote."),
            shape("panel-1", "rect", 96, 216, 528, 336, "transparent", BORDER, 1.5, 28),
            text("bootstrap-tree", 128, 244, 480, 280, esc(
                "README.md", ".github/steps/", "\u00a0" * 4 + "1-step.md ... x-review.md", ".github/workflows/",
                "\u00a0" * 4 + "0-start-exercise.yml", "\u00a0" * 4 + "1-step.yml ... 3-last-step.yml"),
                 21, family=MONO, line=1.55, fx=enter(1, "fade")),
            text("bootstrap-gets", 680, 240, 504, 40, esc("The learner gets"), 30, 700, line=1.2, fx=enter(2)),
            text("bootstrap-gets-body", 680, 300, 504, 130,
                 esc("A clear entry point.", "Issue-driven practice.", "Checks and useful feedback."), 24,
                 color=MUTED, line=1.45, role="body", fx=enter(2)),
            *command("bootstrap", "/bootstrap-github-skills-exercise"),
            *tracker(1),
        ], transition="morph"),
        slide("review", "Feedback is a circuit", [
            art("review", "art-review-circuit"),
            *tracker(2, y=662),
        ], transition="morph"),
        slide("publish", "Publish", [
            *chrome(fx=enter(0, "fade")),
            *heading("publish", "Publish: prepare first. Authorize separately.",
                     "Release-ready content is not permission to publish it."),
            *card("publish", "card-a", "badge-a", 96, 216, 528, 336, "P", "Prepare the handoff",
                  ["Validation evidence.", "PR copy and release notes.", "Launch checklist and known risks."], 1, 30, 22),
            *card("publish", "card-b", "badge-b", 656, 216, 528, 336, "H", "Pause for authorization",
                  ["Confirm the exact owner/repo.", "Confirm visibility.", "Then create, push, or transfer."], 2, 30, 22),
            *command("publish", "/publish-github-skills-exercise"),
            *tracker(3),
        ], transition="morph"),
        slide("recap", "Recap", [
            *chrome(),
            *heading("recap", "One workflow. Many learning experiences.",
                     "Start with expertise. Ship practice, validation, and useful feedback."),
            *card("recap", "card-a", "badge-a", 96, 216, 254, 360, "1", "Repeatable",
                  ["Reuse the same authoring loop across audiences."], 1, 26, 20),
            *card("recap", "card-b", "badge-b", 374, 216, 254, 360, "2", "Observable",
                  ["Check what learners wrote, not file presence."], 2, 26, 20),
            *card("recap", "card-c", "badge-c", 652, 216, 254, 360, "3", "Inclusive",
                  ["Review clarity, navigation, and accessibility."], 3, 26, 20),
            *card("recap", "card-d", "badge-d", 930, 216, 254, 360, "4", "Human-led",
                  ["Approve the decisions and the final release."], 4, 26, 20),
            footnote("recap",
                     "Outline and review generalize to other tutorials. Bootstrap targets the GitHub Skills contract."),
        ], transition="morph"),
        slide("next", "Your next exercise", [
            image("next-texture", 0, 0, 1280, 720, "texture", opacity=0.22, fx=KEN | {"ken": {"dir": "drift", "scale": 1.08, "duration": 24}}),
            *chrome(),
            *heading("next", "Your next exercise starts here."),
            text("next-pitch", 96, 200, 700, 120, esc("Bring an idea.", "Build a hands-on experience."), 40, 700,
                 line=1.3, fx=enter(1)),
            text("next-link", 96, 356, 520, 48, "<u>gh.io/gh-skills-builder</u>", 34, 600, ACCENT, line=1.3,
                 link=PLUGIN_URL, fx=enter(2)),
            text("next-where", 96, 440, 700, 80,
                 esc("Available in Awesome Copilot.", "One orchestrator. Five Agent Skills."), 26, color=MUTED,
                 line=1.4, fx=enter(3)),
            image("next-qr", 864, 176, 320, 320, "qr", fit="contain", link=PLUGIN_URL, fx=enter(2, "fade")),
            text("next-scan", 864, 516, 320, 32, esc("Scan to explore the plugin"), 22, color=MUTED, align="center",
                 line=1.3, fx=enter(3, "fade")),
            footnote("next", "Q&A / What could you turn into an exercise?"),
        ]),
        slide("original-title", "Original title (backup)", [
            image("splash-texture", 0, 0, 1280, 720, "texture", fx=KEN | {"ken": {"dir": "drift", "scale": 1.08, "duration": 24}}),
            svg("splash-mark", 550, 271, 180, 174, "github-mark", fx=enter(0, "fade")),
            text("splash-kicker", 770, 72, 414, 28, esc("LIVE WORKSHOP / 60 MINUTES"), 18, 700, ACCENT, line=1.3,
                 letter=0.5, role="kicker", fx=enter(1)),
            text("splash-title", 770, 148, 414, 180, esc("Turn any idea into a hands-on GitHub exercise"), 44, 700,
                 line=1.2, role="title", fx=enter(1)),
            text("splash-subtitle", 770, 436, 414, 72,
                 esc("Build with gh-skills-builder.", "Live in the GitHub Copilot app."), 24, color=MUTED, line=1.4,
                 role="subtitle", fx=enter(2)),
            text("splash-author", 770, 608, 414, 30, esc("Ari LiVigni"), 24, color=MUTED, line=1.25, fx=enter(3)),
        ], background=PANEL, hidden=True),
    ]


# --- Notes -----------------------------------------------------------------


def notes_text(paragraphs):
    """Flatten sync-notes paragraphs into the plain text Bento shows in speaker view."""
    return "\n".join("".join(text for text, _, _ in runs) for runs in paragraphs).strip()


def build_doc(media, notes, fonts_dir=FONTS):
    slides = drop_morphing_entrances(build_slides())
    numbers = set(range(1, len(slides) + 1))
    if set(notes) != numbers:
        missing, extra = sorted(numbers - set(notes)), sorted(set(notes) - numbers)
        raise BuildError(f"{len(slides)} slides; missing notes for {missing}, unknown slides {extra}")
    for number, s in enumerate(slides, 1):
        s["notes"] = notes_text(notes[number])

    assets, fonts = {}, []
    for family, weight, filename in FONT_FILES:
        key = "font-" + filename.removesuffix(".woff2").lower()
        assets[key] = data_uri("font/woff2", (fonts_dir / filename).read_bytes())
        fonts.append({"family": family, "asset": key, "weight": weight, "style": "normal"})
    for key, (mime, data) in media.items():
        if key == "og-image":
            continue
        assets[key] = data.decode("utf-8") if mime == "image/svg+xml" else data_uri(mime, data)

    return {
        "format": "bento/slides",
        "version": 1,
        "title": "Turn any idea into a Hands-on GitHub Exercise",
        "size": {"width": 1280, "height": 720},
        "theme": {"background": BG, "color": WHITE, "accent": ACCENT, "fontFamily": SANS},
        # Reveal's arrows would sit on the GitHub mark; keys, clickers, swipes, and the site's buttons navigate.
        "present": {"slideNumber": False, "controls": False, "progress": True},
        "meta": {
            "author": "Ari LiVigni",
            "company": "GitHub",
            "subject": "Build a GitHub Skills exercise with the gh-skills-builder plugin",
            "event": "gh-skills-builder workshop",
        },
        "fonts": fonts,
        "slides": slides,
        "assets": assets,
        "modified": "",
    }


def public_doc(doc):
    """The website copy: a read-only player with no speaker notes."""
    doc = copy.deepcopy(doc)
    doc["readonly"] = True
    for s in doc["slides"]:
        s["notes"] = ""
    return doc


# --- Shell splice ----------------------------------------------------------


def serialize(doc):
    # Escaping every "<" keeps "</script>" out of the block (the Bento splice contract).
    return json.dumps(doc, ensure_ascii=False, indent=1).replace("<", "\\u003c")


def read_doc(page):
    match = BLOCK.search(page)
    if not match:
        raise BuildError("no #bento-doc block found in the Bento shell")
    raw = match.group(2).strip()
    return json.loads(raw) if raw else None


def splice(page, doc):
    matches = list(BLOCK.finditer(page))
    if len(matches) != 1:
        raise BuildError(f"expected one #bento-doc block, found {len(matches)}")
    m = matches[0]
    return page[:m.start(2)] + serialize(doc) + page[m.end(2):]


def comparable(doc):
    """Everything but the per-save fields. Images compare byte for byte, so swapped artwork is never missed."""
    if doc is None:
        return None
    return {k: v for k, v in doc.items() if k not in ("docId", "modified")}


def render(path, shell, doc, reshell=False):
    """Return the new page for path, or None when it is already current.

    With reshell, an unchanged document is still moved into the given Bento app.
    """
    existing = read_doc(path.read_text(encoding="utf-8")) if path.exists() else None
    if existing is not None:
        if "docId" in existing:
            doc = {"docId": existing["docId"], **doc}
        if comparable(existing) == comparable(doc) and existing.get("modified"):
            return splice(shell, existing) if reshell else None
    doc["modified"] = datetime.now(timezone.utc).isoformat(timespec="milliseconds").replace("+00:00", "Z")
    return splice(shell, doc)


def load_shell(refresh, shell_path):
    if shell_path:
        return shell_path.read_text(encoding="utf-8")
    if DECK.exists() and not refresh:
        return DECK.read_text(encoding="utf-8")
    print(f"Downloading the Bento app from {SHELL_URL}")
    with urllib.request.urlopen(SHELL_URL, timeout=60) as response:
        return response.read().decode("utf-8")


def main():
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("--check", action="store_true", help="report out-of-date files without writing")
    parser.add_argument("--refresh-shell", action="store_true", help="download the latest Bento app shell")
    parser.add_argument("--shell", type=Path, help="use this local Bento shell file instead")
    args = parser.parse_args()

    try:
        sync_notes = load_sync_notes()
        notes = sync_notes.parse(NOTES_SOURCE.read_text(encoding="utf-8"))
        media = extract_media(PPTX)
        doc = build_doc(media, notes)
        stale = []

        shell = None if args.check else load_shell(args.refresh_shell, args.shell)
        for path, deck_doc in ((DECK, doc), (SITE_DECK, public_doc(doc))):
            rel = path.relative_to(ROOT)
            if args.check:
                existing = read_doc(path.read_text(encoding="utf-8")) if path.exists() else None
                if comparable(existing) != comparable(deck_doc):
                    stale.append(str(rel))
                continue
            page = render(path, shell, deck_doc, reshell=bool(args.refresh_shell or args.shell))
            if page is None:
                print(f"{rel}: up to date")
            else:
                path.parent.mkdir(parents=True, exist_ok=True)
                path.write_text(page, encoding="utf-8")
                print(f"{rel}: written ({len(page.encode('utf-8')) / 1e6:.1f} MB)")

        for name, data in site_images(media).items():
            path = SITE_IMG / name
            rel = path.relative_to(ROOT)
            current = path.read_bytes() if path.exists() else None
            if current == data:
                continue
            if args.check:
                stale.append(str(rel))
            else:
                path.parent.mkdir(parents=True, exist_ok=True)
                path.write_bytes(data)
                print(f"{rel}: written")
    except (BuildError, OSError) as error:
        sys.exit(f"build-bento: {error}")
    except Exception as error:
        if type(error).__name__ == "NotesError":
            sys.exit(f"build-bento: run-of-show.md: {error}")
        raise

    if args.check:
        if stale:
            print("Out of date (run python3 build/build-bento.py):\n  " + "\n  ".join(stale))
            sys.exit(1)
        print("Bento decks and site images are up to date")


if __name__ == "__main__":
    main()
