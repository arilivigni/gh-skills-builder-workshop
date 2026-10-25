import contextlib
import importlib.util
import io
import tempfile
import unittest
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.opc.oxml import serialize_part_xml
from pptx.oxml.xmlchemy import OxmlElement

SPEC = importlib.util.spec_from_file_location(
    "sync_notes", Path(__file__).with_name("sync-notes.py")
)
notes = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(notes)


def source(body, number=1):
    return (
        f"<!-- speaker-notes slide={number} -->\n"
        + "\n".join("> " + line for line in body.splitlines())
        + "\n<!-- /speaker-notes -->"
    )


def visible_text(paragraphs):
    return "\n".join("".join(text for text, _, _ in runs) for runs in paragraphs)


class ParsingTests(unittest.TestCase):
    def test_inline_code_is_literal_even_inside_bold(self):
        self.assertEqual(
            notes.inline_runs("**DO:** Copy `**literal**` and **`<owner>`**.", 1),
            [
                ("DO:", True, False),
                (" Copy ", False, False),
                ("**literal**", False, False),
                (" and ", False, False),
                ("<owner>", True, False),
                (".", False, False),
            ],
        )
        self.assertEqual(
            notes.inline_runs("`[label](target)`", 1),
            [("[label](target)", False, False)],
        )

    def test_unsupported_or_unbalanced_prose_is_rejected(self):
        for text in (
            "Unclosed `code", "Unclosed **bold", "*italic*", "_italic_",
            "~~struck~~", "[label](target)", "<owner>",
        ):
            with self.subTest(text=text), self.assertRaises(notes.NotesError):
                notes.parse(source(text))

    def test_code_blocks_and_repeated_slide_blocks_keep_order(self):
        markdown = source("**DO:** Paste:\n\n```text\n**literal** <owner>\n\n  indented\n```")
        markdown += "\n" + source("**CLOCK:** Stop.", 1)
        parsed = notes.parse(markdown)
        self.assertEqual(
            visible_text(parsed[1]),
            "DO: Paste:\n\n**literal** <owner>\n\n  indented\n\nCLOCK: Stop.",
        )
        self.assertEqual(parsed[1][2], [("**literal** <owner>", False, True)])

    def test_invalid_markers_and_fences_are_rejected(self):
        for markdown in (
            "<!-- /speaker-notes -->",
            "<!-- speaker-notes slide=1 -->\n> Unclosed",
            "<!-- speaker-notes slide=1 -->\n<!-- speaker-notes slide=2 -->",
            source("```text\nUnclosed fence"),
            "<!-- speaker-notes slide=1 -->\nNot quoted\n<!-- /speaker-notes -->",
        ):
            with self.subTest(markdown=markdown), self.assertRaises(notes.NotesError):
                notes.parse(markdown)


class SynchronizationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.deck = Path(self.temp.name) / "test.pptx"
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.notes_slide.notes_text_frame.text = "Original notes"
        presentation.save(self.deck)
        self.expected = notes.parse(source("**SAY:** Hello.\n\n**DO:** Paste:\n\n```text\ncommand <owner>\n```"))

    def sync(self, check=False, expected=None):
        with contextlib.redirect_stdout(io.StringIO()):
            return notes.sync(self.deck, self.expected if expected is None else expected, check)

    def parts(self):
        with zipfile.ZipFile(self.deck) as archive:
            return {name: archive.read(name) for name in archive.namelist()}

    def test_only_notes_change_and_repeated_sync_is_a_noop(self):
        before = self.parts()
        original = self.deck.read_bytes()
        self.assertFalse(self.sync(check=True))
        self.assertEqual(self.deck.read_bytes(), original)
        self.assertFalse(self.sync())
        after = self.parts()
        self.assertEqual(before.keys(), after.keys())
        changed = [name for name in before if before[name] != after[name]]
        self.assertEqual(changed, ["ppt/notesSlides/notesSlide1.xml"])
        frame = Presentation(self.deck).slides[0].notes_slide.notes_text_frame
        self.assertEqual(frame.text, visible_text(self.expected[1]))
        self.assertEqual(notes.read_notes(frame), notes.normalize(self.expected[1]))
        updated = self.deck.read_bytes()
        self.assertTrue(self.sync(check=True))
        self.assertTrue(self.sync())
        self.assertEqual(self.deck.read_bytes(), updated)

    def test_powerpoint_edits_are_detected_and_repaired(self):
        for edit in ("text", "bold", "font", "soft-break", "field"):
            with self.subTest(edit=edit):
                self.sync()
                presentation = Presentation(self.deck)
                part = presentation.slides[0].notes_slide.part
                frame = presentation.slides[0].notes_slide.notes_text_frame
                paragraph = frame.paragraphs[0]
                if edit == "text":
                    paragraph.runs[1].text += " Changed."
                elif edit == "bold":
                    paragraph.runs[0].font.bold = False
                elif edit == "font":
                    frame.paragraphs[-1].runs[0].font.name = "Arial"
                elif edit == "soft-break":
                    paragraph.add_line_break()
                else:
                    field = OxmlElement("a:fld")
                    field.set("id", "{00000000-0000-0000-0000-000000000001}")
                    field.set("type", "slidenum")
                    text = OxmlElement("a:t")
                    text.text = "1"
                    field.append(text)
                    paragraph._p.append(field)
                notes.replace_parts(self.deck, {
                    str(part.partname).lstrip("/"): serialize_part_xml(part._element)
                })
                changed = self.deck.read_bytes()
                self.assertFalse(self.sync(check=True))
                self.assertEqual(self.deck.read_bytes(), changed)
                self.assertFalse(self.sync())
                self.assertTrue(self.sync(check=True))

    def test_missing_or_extra_slides_fail_without_writing(self):
        original = self.deck.read_bytes()
        for expected in ({}, {1: self.expected[1], 2: self.expected[1]}):
            with self.subTest(expected=expected), self.assertRaises(notes.NotesError):
                self.sync(expected=expected)
            self.assertEqual(self.deck.read_bytes(), original)

    def test_committed_decks_match_the_run_of_show(self):
        expected = notes.parse(notes.SOURCE.read_text(encoding="utf-8"))
        for deck in notes.DECKS:
            with self.subTest(deck=deck.name):
                presentation = Presentation(deck)
                self.assertEqual(set(expected), set(range(1, len(presentation.slides) + 1)))
                for number, slide in enumerate(presentation.slides, 1):
                    self.assertTrue(slide.has_notes_slide)
                    frame = slide.notes_slide.notes_text_frame
                    self.assertEqual(frame.text, visible_text(expected[number]))
                    self.assertEqual(notes.read_notes(frame), notes.normalize(expected[number]))


if __name__ == "__main__":
    unittest.main()
