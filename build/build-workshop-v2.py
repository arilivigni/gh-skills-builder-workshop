#!/usr/bin/env python3
"""Rebuild the workshop on the official GitHub template, with plugin onboarding."""
from pathlib import Path
from io import BytesIO
import importlib.util
from PIL import Image

HERE = Path(__file__).resolve().parent
spec = importlib.util.spec_from_file_location("template_builder", HERE / "build-from-template.py")
base = importlib.util.module_from_spec(spec)
spec.loader.exec_module(base)

from pptx import Presentation
from pptx.util import Inches
from pptx.enum.text import PP_ALIGN

ROOT = HERE.parent
VISUALS = ROOT / "visuals"
QR = VISUALS / "gh-qr-gh-skills-builder.png"
LINK = "https://gh.io/gh-skills-builder"
INSTALL = "copilot plugin install gh-skills-builder@awesome-copilot"
GEMINI = {
    "01-agent-skill-map": {
        "light": "Gemini_Generated_Image_ggbn2sggbn2sggbn.jpeg",
        "dark": "Gemini_Generated_Image_82nv0i82nv0i82nv.jpeg",
    },
    "02-approval-gated-lifecycle": {
        "light": "Gemini_Generated_Image_zg44zwzg44zwzg44.jpeg",
        "dark": "Gemini_Generated_Image_tz1drutz1drutz1d.jpeg",
    },
    "03-review-fix-circuit": {
        "light": "Gemini_Generated_Image_5go8r55go8r55go8.jpeg",
        "dark": "Gemini_Generated_Image_1uodcb1uodcb1uod.jpeg",
    },
}
txt, box, note = base.add_text, base.add_box, base.set_notes
base.FONT_MONO = "Hubot Sans"


def link_text(slide, x, y, w, h, text, url, **kwargs):
    shape = txt(slide, x, y, w, h, text, **kwargs)
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            r.hyperlink.address = url
    return shape


def build(theme):
    t = base.THEMES[theme]
    prs = Presentation(str(base.TEMPLATE_PPTX))
    layouts = {
        key: base.layout_by_name(prs, name, t["master"])
        for key, name in [
            ("splash", "Splash page"), ("agenda", "Agenda"),
            ("content", "Title and Content"), ("four", "Four statements columns"),
        ]
    }
    old_ids = list(prs.slides._sldIdLst)

    def new(layout="content"):
        s = prs.slides.add_slide(layouts[layout])
        for ph in list(s.placeholders):
            ph._element.getparent().remove(ph._element)
        return s

    def text(s, x, y, w, h, text, size=20, color=None, bold=False, **kw):
        return txt(s, x, y, w, h, text, size=size, color=color or t["text"],
                   bold=bold, font=base.FONT_DISPLAY if bold else base.FONT_BODY,
                   margin=0, fit=False, **kw)

    def heading(s, title, sub=None):
        text(s, .65, .48, 12, .74, title, 34, bold=True)
        if sub:
            text(s, .65, 1.35, 11.9, .8, sub, 21, t["body"])

    def footer(s, label):
        text(s, .65, 6.65, 11.9, .25, label, 11, t["muted"])

    def card(s, x, y, w, h, title, body, mark=None):
        box(s, x, y, w, h, t["panel"], t["outline"])
        top = y + .25
        if mark:
            base.add_badge(s, x + .24, top, .52, mark, t, t["green_deep"])
            top += .8
        text(s, x + .25, top, w - .5, .8, title, 24, bold=True)
        text(s, x + .25, top + .95, w - .5, h - (top - y) - 1.15,
             body, 19, t["body"])

    def art(filename, notes):
        s = new()
        image = BytesIO()
        with Image.open(VISUALS / GEMINI[filename][theme]) as source:
            source.save(image, format="PNG")
        image.seek(0)
        picture = s.shapes.add_picture(image, 0, 0, width=prs.slide_width)
        picture.top = int((prs.slide_height - picture.height) / 2)
        picture._element.xpath("./p:nvPicPr/p:cNvPr")[0].set(
            "descr", {
                "01-agent-skill-map": "One orchestrator agent loads five skills: orchestration, outline, bootstrap, review, and publish.",
                "02-approval-gated-lifecycle": "Outline, bootstrap, review, and publish, with human approval at every transition.",
                "03-review-fix-circuit": "Build, report-only review, human decision, approved fixes, and re-review before publication.",
            }[filename])
        note(s, notes)
        return s

    # 1: Keep the template splash artwork, including its centered GitHub logo.
    s = new("splash")
    text(s, 8.02, .76, 4.65, .4, "LIVE WORKSHOP / 60 MINUTES", 14, t["green"], True)
    text(s, 8.02, 1.52, 4.65, 2.8,
         "Turn any idea into a hands-on GitHub exercise", 36, bold=True)
    text(s, 8.02, 4.55, 4.65, 1.1,
         "Build with gh-skills-builder.\nLive in the GitHub Copilot app.", 20, t["body"])
    text(s, 8.02, 6.35, 4.65, .4, "Ari LiVigni", 18, t["muted"])
    note(s, "0:00-0:01 (slides 1-2). Welcome. In 60 minutes we turn one small idea into a "
         "self-paced GitHub Skills exercise and see all four builder phases. The exercise: "
         "Cleared for Takeoff, a 15-20-minute Copilot quick-start inside Tailspin Toys with "
         "three graded steps. Most of this session is a live walkthrough; the slides set the stage. "
         "Content ends by 0:45 and Q&A is protected from 0:50.")

    # 2: Agenda is deliberately second, on the template's Agenda layout.
    s = new("agenda")
    text(s, .6, 3.12, 3.2, 1.45, "Our 60-minute\nflight plan", 30, bold=True)
    text(s, .6, 5.14, 3.2, 1.1, "Mostly live.\nCheckpoints protect the demo.", 21, t["body"])
    agenda = [
        ("00-06", "Get oriented", "Plugins, installation, and the quick-start scenario"),
        ("06-16", "Outline + hand-off", "Three learner steps, a name, and a private repo"),
        ("16-27", "Bootstrap", "From outline to files, plus one fail-to-pass check"),
        ("27-35", "Review", "Report findings, approve one fix, review again"),
        ("35-41", "Publish readiness", "Evidence and a launch plan, without publishing"),
        ("41-60", "Recap + Q&A", "Takeaways by 0:45, then a buffer and Q&A by 0:50"),
    ]
    for i, (time, title, desc) in enumerate(agenda):
        y = .57 + i * 1.02
        text(s, 4.83, y, 1.25, .4, time, 19, t["green"], True)
        text(s, 6.3, y, 6.25, .4, title, 23, bold=True)
        text(s, 6.3, y + .47, 6.25, .5, desc, 17, t["body"])
    note(s, "0:00-0:01, with slide 1. Time budget: 6 minutes orientation (0:00-0:06), 6 outline "
         "(0:06-0:12), 4 hand-off (0:12-0:16), 11 bootstrap (0:16-0:27), 8 review (0:27-0:35), "
         "6 publish readiness (0:35-0:41), 4 recap (0:41-0:45), a 5-minute flex buffer "
         "(0:45-0:50), and 10 minutes of protected Q&A (0:50-1:00). Hard stops: saved outline "
         "at 0:12, rehearsal 00-idea at 0:16, 02-bootstrapped at 0:19, saved review at 0:30, "
         "no more fix rounds at 0:35, saved readiness at 0:38. If the buffer isn't needed, "
         "open Q&A at 0:45. Installation is a 60-second demonstration at most; attendee "
         "troubleshooting happens after the session.")

    # 3: Plugin definition and a visual package, without conflating package and agent.
    s = new()
    heading(s, "What is a plugin?",
            "An installable toolkit that adds reusable capabilities to GitHub Copilot.")
    box(s, .7, 2.55, 5.25, 3.65, t["card"], t["outline"])
    base.add_badge(s, 1.0, 2.9, .75, "+", t, t["green_deep"])
    text(s, 2.0, 2.98, 3.6, .55, "gh-skills-builder", 27, bold=True)
    text(s, 1.05, 4.03, 4.55, 1.45,
         "Install once.\nReuse the workflow.", 27, t["green"], True)
    card(s, 6.42, 2.55, 2.95, 3.65, "One agent",
         "Coordinates the work and asks for your approval.", "A")
    card(s, 9.7, 2.55, 2.95, 3.65, "Five skills",
         "Reusable workflows and capabilities the agent loads.", "S")
    footer(s, "A plugin packages customizations. It is not a separate AI model.")
    note(s, "0:01-0:03 (slides 3-4). Slide-ready explanation: A plugin is an installable toolkit that adds reusable "
         "capabilities to GitHub Copilot. Instead of copying prompts and customization files "
         "by hand, install a bundle and use it in your work. gh-skills-builder packages "
         "one orchestrator agent and five Agent Skills for the exercise lifecycle. "
         "Agent SKILLS are automated workflows/capabilities for an agent, not five separate agents. "
         "Skills guide agent behavior; they do not guarantee correctness or replace approvals. "
         "Other plugins can bundle additional kinds of supported customizations. "
         "Source: https://github.com/github/awesome-copilot/blob/main/docs/README.plugins.md")

    # 4: Marketplace is the primary install path. Keep the provided QR untouched.
    s = new()
    heading(s, "Add the toolkit. Start building.",
            "gh-skills-builder is available in the Awesome Copilot marketplace.")
    for i, (title, body) in enumerate([
        ("Find it", "Open Plugins in the Copilot app.\nFind gh-skills-builder."),
        ("Review + install", "Review the source and capabilities.\nInstall and enable for your session."),
        ("Use it", "Choose the orchestrator agent.\nAsk it to run the exercise lifecycle."),
    ]):
        y = 2.4 + i * 1.12
        base.add_badge(s, .76, y, .48, str(i + 1), t, t["green_deep"])
        text(s, 1.48, y, 6.7, .43, title, 23, bold=True)
        text(s, 1.48, y + .46, 6.65, .65, body, 17, t["body"])
    s.shapes.add_picture(str(QR), Inches(9.25), Inches(2.35),
                         width=Inches(2.8), height=Inches(2.8))
    link_text(s, 8.55, 5.43, 4.2, .55, "gh.io/gh-skills-builder", LINK,
              size=20, color=t["green"], bold=True, font=base.FONT_BODY,
              align=PP_ALIGN.CENTER, margin=0)
    box(s, .75, 6.22, 11.85, .52, t["panel"], t["outline"])
    text(s, .96, 6.32, 11.4, .33, "CLI: " + INSTALL, 16, t["text"])
    note(s, "0:01-0:03, with slide 3. Cap the install demo at 60 seconds, or skip it and leave the QR up. "
         "Demo installing from Plugins in the Copilot app, then enable it for this session. "
         "Use the QR/link to open the Awesome Copilot plugin page; the QR does not auto-install. "
         "CLI command: " + INSTALL + ". If an older/custom CLI says the marketplace is unknown, run "
         "'copilot plugin marketplace add github/awesome-copilot' once, then retry. "
         "Check installation with 'copilot plugin list'. In an interactive CLI session: "
         "'/plugin install gh-skills-builder@awesome-copilot'. VS Code alternative: "
         "Extensions search '@agentPlugins', or Command Palette > Chat: Plugins. "
         "Review external source before installing. The marketplace listing is pinned to a release; "
         "new local or unreleased skill changes are not automatically available from that listing. "
         "Sources: https://github.com/github/awesome-copilot#install-a-plugin ; "
         "https://awesome-copilot.github.com/plugin/gh-skills-builder/")

    # 5: Clearly distinguish the authoring demo from the learner's feature work.
    s = new()
    heading(s, "Today's scenario: Tailspin Toys",
            "Build a 15–20-minute Copilot quick-start inside Tailspin Toys.")
    card(s, .75, 2.65, 5.45, 3.35, "Start with a real app",
         "Tailspin Toys as MIT starter code. Learners customize Copilot, not app features.", "1")
    card(s, 7.12, 2.65, 5.45, 3.35, "Teach three habits",
         "Repository instructions.\nConsistent quality reports.\nA report-only QA agent.", "2")
    base.add_badge(s, 6.37, 3.85, .48, ">", t, t["green_deep"])
    text(s, .85, 6.35, 11.6, .42,
         "Working title: Cleared for Takeoff.", 24, t["green"], True)
    note(s, "0:03-0:04. Source scenario: the Dev Days Copilot app lab "
         "(https://github-samples.github.io/copilot-workshops/real-world-development/app/). "
         "We build a small slice of it, not the whole lab. Starter: "
         "https://github.com/github-samples/tailspin-toys (MIT). The learner makes three changes: "
         "a TSDoc section in .github/copilot-instructions.md, an evidence-based reporting contract "
         "in the existing quality-checks skill, and a report-only .github/agents/qa.agent.md. "
         "No app feature, filtering, npm install, database, or browser tooling. "
         "We do not copy the Tailspin template: handoff.sh imports the app as one MIT starter commit "
         "and drops the upstream issue sources and their self-deleting seeding workflow, so Step 0 "
         "opens only the exercise issue. "
         "Use candidate name copilot-app-cleared-for-takeoff; learner copies use skills-<repo>. "
         "Ask for names explicitly so the prompt works with the currently indexed release. "
         "Actions checks what learners wrote; it cannot observe local app modes or prove "
         "future model behavior.")

    # 6: Human role, not a claim that autonomous generation replaces judgment.
    s = new()
    heading(s, "From builder to architect",
            "The agent does the heavy lifting. You own the learning experience.")
    card(s, .75, 2.55, 5.65, 3.55, "Let the agent build",
         "Draft step content.\nScaffold repository files.\nWire validation and feedback.", "A")
    card(s, 6.92, 2.55, 5.65, 3.55, "Keep the human decisions",
         "Set learner outcomes.\nApprove the outline and fixes.\nAuthorize publication.", "H")
    footer(s, "Same pattern for educators, maintainers, event organizers, DevRel, startups, and enterprises.")
    note(s, "0:04-0:05. Quick show of hands, 10 seconds at most; skip it if behind: who teaches, "
         "onboards contributors, runs events, or trains teams? The workflow is reusable across "
         "these audiences. Today's anchor is onboarding a new Tailspin teammate to the team's "
         "Copilot conventions. Your role shifts toward architect and reviewer, but you remain "
         "responsible for the result.")

    # 7-8: Use the user's Gemini artwork in the matching light/dark variant.
    art("01-agent-skill-map",
        "0:05-0:06 (slides 7-8). One orchestrator, five skills. The orchestration skill runs the full lifecycle; "
        "the four phase skills own outline, bootstrap, review, and publish requirements. "
        "Agent SKILLS are automated workflows/capabilities for an agent, not five separate agents. "
        "Select @github-skills-orchestrator and use /orchestrate-github-skills-exercise "
        "or call a phase skill directly.")
    art("02-approval-gated-lifecycle",
        "0:05-0:06, with slide 7. Four phases, explicit human gates. Approve the outline before generation; approve "
        "the build before review; choose fixes or sign off; authorize actual publication separately. "
        "A diagrammed approval is behavioral guidance, not an enforced permission boundary. "
        "Show Prompt 1 from run-of-show.md next, in a quick chat: we start with only an idea, no repo. "
        "We're now at minute 6.")

    # 9: Outline shows the idea-to-contract transformation.
    s = new()
    heading(s, "Outline: decide what success looks like",
            "Input: a Tailspin idea. Output: an approved three-step journey.")
    card(s, .75, 2.55, 3.8, 3.45, "Outcome",
         "Three small Copilot customizations, done in 15–20 minutes.", "1")
    card(s, 4.77, 2.55, 3.8, 3.45, "Identity",
         "Choose an exercise title and an eye-catching repository name.", "2")
    card(s, 8.79, 2.55, 3.8, 3.45, "Validation",
         "For every graded step: what can a workflow actually check?", "3")
    text(s, .8, 6.3, 11.8, .42, "/create-github-skills-outline", 24, t["green"], True)
    note(s, "0:06-0:12, keep fully live. Quick chat: Interactive mode, Auto > Intelligence. Paste Prompt 1. "
         "Narrate one Theory/Activity pair and its validation signal. Ask: what would prove the learner "
         "changed the reporting contract, rather than merely created a file? A bare file-exists check is "
         "not evidence. Show three title/name candidates; cap the vote at 30 seconds and use the rehearsed "
         "default, copilot-app-cleared-for-takeoff, if needed. Gate 1, sent once: approve the outline and "
         "identity. Nothing is created yet. Check the returned single-block artifact starts with the title, "
         "Repository name, and Learner copy, then copy the whole block. "
         "Hard stop 0:12: use the saved fallback/exercise-outline.md and say it is from rehearsal. "
         "Hand-off, 0:12-0:16, stay on this slide. Say: Gate 1 approved the idea and created nothing; "
         "the Hand-off Authorization is a separate approval for one private authoring repo; publishing is a "
         "later gate. In the same quick chat, send the Hand-off Authorization from run-of-show.md and approve the "
         "shell command out loud (the terminal is the fallback). handoff.sh <you> <approved-name> creates the "
         "private repo, imports Tailspin Toys as one MIT starter commit without the upstream issue sources, "
         "issue-seeding and triage workflows, or CODEOWNERS, and pushes. In the Copilot app, Add project on "
         "the folder it printed, then New session on main, Interactive, Auto > Balance. Paste Prompt 2 there "
         "with the whole outline artifact; the new session cannot see the quick-chat history. "
         "Hard stop 0:16: switch to the pre-added rehearsal project at checkpoint 00-idea.")

    # 10: A readable file tree beside learner-visible outcomes.
    s = new()
    heading(s, "Bootstrap: make the outline runnable",
            "Generate exercise files in the current repository, not a new GitHub remote.")
    box(s, .75, 2.5, 6.2, 3.55, t["panel"], t["outline"])
    text(s, 1.05, 2.85, 5.6, 2.85,
         "README.md\n.github/steps/\n    1-step.md ... x-review.md\n.github/workflows/\n    0-start-exercise.yml\n    1-step.yml ... 3-last-step.yml", 21)
    text(s, 7.45, 2.72, 5.05, .5, "The learner gets", 27, bold=True)
    text(s, 7.45, 3.55, 5.0, 1.85,
         "A clear entry point.\nIssue-driven practice.\nChecks and useful feedback.", 24, t["body"])
    text(s, .8, 6.3, 11.8, .42, "/bootstrap-github-skills-exercise", 23, t["green"], True)
    note(s, "0:16-0:27. Project session on the new repo: Interactive, Auto > Balance. Paste Prompt 2, "
         "replacing only its placeholder with the whole outline artifact; its header carries the approved "
         "title and repository name, and the new session cannot see the quick-chat history. Its first "
         "action saves the outline verbatim to .github/exercise-outline.md; show the file when it appears. "
         "Then show the README and Copy Exercise badge, the three step files, completion material, and "
         "workflow chaining. Learner work stays unsolved; keep the Tailspin MIT LICENSE and credit. "
         "Step 0 opens only the exercise instructions issue; there is no other issue seeding. "
         "Cap the wait at three minutes. Hard stop 0:19: if the build isn't ready to inspect, switch to the "
         "rehearsal 02-bootstrapped state and say so. Then one fail-to-pass example on step 1, in a separate "
         "disposable learner copy: the untouched starter fails with an actionable message; add the TSDoc "
         "section; the check passes and the next step opens. Checks validate what learners wrote, not file "
         "existence or future Copilot behavior. Use saved evidence or a clip if Actions is slow. "
         "Bootstrap writes into this private authoring repo; publish creates the public learner-facing "
         "template only after separate authorization. Gate 2 by 0:27: accept the build only if its evidence "
         "supports it, then move to the report-only review.")

    # 11: The review circuit is more useful than a text-only divider.
    art("03-review-fix-circuit",
        "0:27-0:35. /review-github-skills-exercise. Same session: Interactive, Auto > Balance. Paste Prompt 3. "
        "Three questions: does the untouched starter fail, does correct learner work pass, and is the "
        "feedback clear and safe? Review reports findings without editing files. Hard stop 0:30: use the "
        "saved review if the live one is unfinished. Pick one real, supported finding, such as a weak grader, "
        "a workflow collision, or a learner accessibility issue; don't manufacture a defect. Approve only "
        "that fix with Prompt 3b, inspect the change and its checks, accept the fix round, then review "
        "again. If there are no live findings, show the labeled rehearsal finding and its before/after. "
        "Hard stop 0:35: start no more rounds. Self-review is not independent review; report-only "
        "guidance does not remove edit permissions. If stuck on the same finding after two rounds, stop "
        "and escalate. Give Gate 3 only when there are no known release blockers.")

    # 12: Put the publication boundary in the visual hierarchy.
    s = new()
    heading(s, "Publish: prepare first. Authorize separately.",
            "Release-ready content is not permission to publish it.")
    card(s, .75, 2.55, 5.65, 3.45, "Prepare the handoff",
         "Validation evidence.\nPR copy and release notes.\nLaunch checklist and known risks.", "P")
    card(s, 6.92, 2.55, 5.65, 3.45, "Pause for authorization",
         "Confirm the exact owner/repo.\nConfirm visibility.\nThen create, push, or transfer.", "H")
    text(s, .8, 6.3, 11.8, .42, "/publish-github-skills-exercise", 24, t["green"], True)
    note(s, "0:35-0:41. Same session: Interactive, Auto > Balance. Never Autopilot here: the authorization "
         "gate must stop for you. Paste Prompt 4, or show the saved 04-publish-ready materials. Keep the "
         "output short: observed validation evidence, untested risks, PR description, release notes, and "
         "launch checklist. Local checks alone don't prove a fresh GitHub copy works. Send Prompt 4b with "
         "<demo-owner>/copilot-app-cleared-for-takeoff to show the plan and the publication boundary. "
         "Do not authorize it. The earlier private repo creation was not permission to publish, transfer, "
         "change visibility, or create a public template. Publish can create a GitHub repository, but needs "
         "explicit owner/repo and visibility authorization. Reconcile the README badge with the approved "
         "target. Hard stop 0:38: show the saved readiness report and planned commands. End at 0:41.")

    # 13: Four takeaways with visual anchors.
    s = new("four")
    heading(s, "One workflow. Many learning experiences.",
            "Start with expertise. Ship practice, validation, and useful feedback.")
    for i, (title, body) in enumerate([
        ("Repeatable", "Reuse the same authoring loop across audiences."),
        ("Observable", "Check what learners wrote, not file presence."),
        ("Inclusive", "Review clarity, navigation, and accessibility."),
        ("Human-led", "Approve the decisions and the final release."),
    ]):
        card(s, .75 + i * 3.04, 2.7, 2.75, 3.4, title, body, str(i + 1))
    footer(s, "Outline and review generalize to other tutorials. Bootstrap targets the GitHub Skills contract.")
    note(s, "0:41-0:45. Show the /orchestrate-github-skills-exercise recap prompt without running it: that is "
         "the fifth skill, and it sequences the same lifecycle; it is not a second agent. Mention filtering "
         "and tests, local Plan/Autopilot practice, PR and Agent Merge, MCP, and canvases as extensions for a "
         "longer exercise; don't demo them. Recap for an educator (a classroom convention) and an enterprise "
         "(a team onboarding policy): same lifecycle, different content. SemVer and plugin CI describe plugin "
         "release hygiene, not a guarantee that a generated exercise works; run and review it with a learner. "
         "Closing line: Keep the learner task small. Make success observable. Let the agent draft; keep humans "
         "at the decisions. Content ends at 0:45. The 0:45-0:50 buffer only absorbs overruns; if unused, "
         "move to slide 14 and open Q&A.")

    # 14: QR at a scan-friendly size, not hidden in a tiny footer.
    s = new()
    heading(s, "Your next exercise starts here.")
    text(s, .8, 2.2, 7.2, 1.4, "Bring an idea.\nBuild a hands-on experience.", 34, bold=True)
    link_text(s, .8, 4.04, 7.4, .7, "gh.io/gh-skills-builder", LINK,
              size=31, color=t["green"], bold=True, font=base.FONT_BODY, margin=0)
    text(s, .8, 5.07, 7.25, .9,
         "Available in Awesome Copilot.\nOne orchestrator. Five Agent Skills.", 23, t["body"])
    s.shapes.add_picture(str(QR), Inches(9.0), Inches(2.0),
                         width=Inches(3.55), height=Inches(3.55))
    text(s, 9.0, 5.88, 3.55, .55, "Scan to explore the plugin", 18, t["muted"],
         align=PP_ALIGN.CENTER)
    footer(s, "Q&A / What could you turn into an exercise?")
    note(s, "0:50-1:00, or from 0:45 if the buffer was unused. Leave this QR on screen. The link resolves "
         "to the Awesome Copilot gh-skills-builder page, not directly to the source repository. Install "
         "command: " + INSTALL + ". Questions to seed: What can Actions observe? Why is review report-only? "
         "What authorization is required to publish? What would your audience practice? Don't troubleshoot "
         "individual installations or start another workflow.")

    for number, slide_obj in enumerate(list(prs.slides)[len(old_ids):], start=1):
        if number > 1:
            text(slide_obj, 12.15, 7.03, .55, .24, str(number), 11, t["muted"],
                 align=PP_ALIGN.RIGHT, margin=0, fit=False)

    for sid in old_ids:
        prs.part.drop_rel(sid.rId)
        prs.slides._sldIdLst.remove(sid)
    prs.core_properties.title = "Turn Any Idea into a Hands-on GitHub Exercise"
    prs.core_properties.author = "Ari LiVigni"
    prs.core_properties.subject = "60-minute workshop / Tailspin quick-start / gh-skills-builder"
    out = ROOT / f"gh-skills-builder-workshop-{theme}.pptx"
    prs.save(str(out))
    base.prune_template_package(out, t["master"])
    return out


if __name__ == "__main__":
    base.convert_potx_to_pptx()
    for theme in ("light", "dark"):
        print(build(theme))
